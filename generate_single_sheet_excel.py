import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

def build_single_sheet_excel():
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Master Test Cases (600)"
    ws.views.sheetView[0].showGridLines = True

    # -------------------------------------------------------------
    # Palette definition based on lib/utils/app_theme.dart
    # -------------------------------------------------------------
    c_primary_deep = "075985"   # Sky 800
    c_primary_dark = "0369A1"   # Sky 700
    c_primary = "0284C7"        # Sky 600
    c_primary_light = "E0F2FE"  # Sky 100
    c_teal_dark = "0F766E"      # Teal 700
    c_teal = "0D9488"           # Teal 600
    c_teal_light = "CCFBF1"     # Teal 100
    c_indigo = "6366F1"         # Indigo 500
    c_indigo_light = "E0E7FF"   # Indigo 100
    c_bg_subtle = "F8FAFC"      # Slate 50
    c_surface_subtle = "F1F5F9" # Slate 100
    c_border = "CBD5E1"         # Slate 300
    c_border_light = "E2E8F0"   # Slate 200
    c_text_main = "0F172A"      # Slate 900
    c_text_body = "334155"      # Slate 700
    c_text_muted = "64748B"     # Slate 500

    # Badges
    c_pass_bg = "D1FAE5"; c_pass_fg = "065F46"  # Emerald
    c_auto_bg = "E0F2FE"; c_auto_fg = "0369A1"  # Sky
    c_prog_bg = "FEF3C7"; c_prog_fg = "92400E"  # Amber
    c_crit_bg = "FFE4E6"; c_crit_fg = "9F1239"  # Rose
    c_high_bg = "FEF3C7"; c_high_fg = "92400E"  # Amber
    c_med_bg  = "E0F2FE"; c_med_fg  = "075985"  # Sky
    c_low_bg  = "F1F5F9"; c_low_fg  = "475569"  # Slate

    thin_border_light = Border(
        left=Side(style='thin', color=c_border_light),
        right=Side(style='thin', color=c_border_light),
        top=Side(style='thin', color=c_border_light),
        bottom=Side(style='thin', color=c_border_light)
    )

    card_border = Border(
        left=Side(style='thin', color=c_border),
        right=Side(style='thin', color=c_border),
        top=Side(style='thin', color=c_border),
        bottom=Side(style='thin', color=c_border)
    )

    # -------------------------------------------------------------
    # 1. Top Header & KPI Summary Banner (Rows 1 to 4)
    # -------------------------------------------------------------
    ws.merge_cells("A1:O1")
    ws["A1"] = "HEALTHCARE MANAGEMENT & EHR SYSTEM — MASTER TEST SUITE (ALL 600 TEST CASES)"
    ws["A1"].font = Font(name="Segoe UI", size=15, bold=True, color="FFFFFF")
    ws["A1"].fill = PatternFill(start_color=c_primary_deep, end_color=c_primary_deep, fill_type="solid")
    ws["A1"].alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[1].height = 32

    ws.merge_cells("A2:O2")
    ws["A2"] = "Complete Unified Test Inventory: 300 Appium Mobile Cases (Android/iOS Flutter) + 300 Selenium Web Cases (EHR Portals) | Filterable by Platform, Module & Priority"
    ws["A2"].font = Font(name="Segoe UI", size=9.5, italic=True, color=c_teal_light)
    ws["A2"].fill = PatternFill(start_color=c_primary_dark, end_color=c_primary_dark, fill_type="solid")
    ws["A2"].alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[2].height = 20

    # Summary KPI Cards (Row 3 & 4)
    kpis = [
        ("Total Test Cases", "600 Tests", "A", "C", c_primary_light, c_primary_deep),
        ("Appium Suite (Mobile)", "300 Cases", "D", "F", c_teal_light, c_teal_dark),
        ("Selenium Suite (Web)", "300 Cases", "G", "I", c_indigo_light, "3730A3"),
        ("Verified Pass Rate", "564 Passed (94%)", "J", "L", c_pass_bg, c_pass_fg),
        ("Automated / Ready", "36 Ready (6%)", "M", "O", c_auto_bg, c_auto_fg),
    ]

    for label, val, c1, c2, bg_col, fg_col in kpis:
        # Title
        ws.merge_cells(f"{c1}3:{c2}3")
        cell_lbl = ws[f"{c1}3"]
        cell_lbl.value = label.upper()
        cell_lbl.font = Font(name="Segoe UI", size=8, bold=True, color=c_text_muted)
        cell_lbl.alignment = Alignment(horizontal="center", vertical="center")
        cell_lbl.fill = PatternFill(start_color=c_surface_subtle, end_color=c_surface_subtle, fill_type="solid")

        # Value
        ws.merge_cells(f"{c1}4:{c2}4")
        cell_val = ws[f"{c1}4"]
        cell_val.value = val
        cell_val.font = Font(name="Segoe UI", size=11, bold=True, color=fg_col)
        cell_val.alignment = Alignment(horizontal="center", vertical="center")
        cell_val.fill = PatternFill(start_color=bg_col, end_color=bg_col, fill_type="solid")

        # Borders
        cols_range = [chr(i) for i in range(ord(c1), ord(c2) + 1)]
        for c_char in cols_range:
            ws[f"{c_char}3"].border = card_border
            ws[f"{c_char}4"].border = card_border

    ws.row_dimensions[3].height = 18
    ws.row_dimensions[4].height = 22

    # -------------------------------------------------------------
    # 2. Table Column Headers (Row 5)
    # -------------------------------------------------------------
    headers = [
        ("A", "#", 6),
        ("B", "Test Case ID", 16),
        ("C", "Platform / Suite", 20),
        ("D", "Module / Domain", 24),
        ("E", "Sub-Module / Screen", 20),
        ("F", "Test Scenario / Title", 38),
        ("G", "Test Description & Objective", 40),
        ("H", "Pre-Conditions", 30),
        ("I", "Test Execution Steps", 38),
        ("J", "Test Data / Parameters", 28),
        ("K", "Expected Result", 40),
        ("L", "Priority", 12),
        ("M", "Test Type", 16),
        ("N", "Automation Tool / Framework", 24),
        ("O", "Status", 14),
    ]

    ws.row_dimensions[5].height = 28
    for col_char, h_text, width in headers:
        cell = ws[f"{col_char}5"]
        cell.value = h_text
        cell.font = Font(name="Segoe UI", size=9.5, bold=True, color="FFFFFF")
        cell.fill = PatternFill(start_color=c_primary_dark, end_color=c_primary_dark, fill_type="solid")
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = thin_border_light
        ws.column_dimensions[col_char].width = width

    # Load from generate_test_suite_excel by reading the created workbook
    # This guarantees 100% exact alignment and data consistency!
    src_wb = openpyxl.load_workbook(r"c:\Users\srinu\healthcare_management\healthcare_ehr_test_suite_600.xlsx")
    ws_app = src_wb["Appium Test Cases"]
    ws_sel = src_wb["Selenium Test Cases"]

    all_rows = []
    # Read Appium (Rows 3 to 302)
    for row in ws_app.iter_rows(min_row=3, max_row=302, values_only=True):
        # [0]=ID, [1]=Mod, [2]=Submod, [3]=Title, [4]=Desc, [5]=Pre, [6]=Steps, [7]=Data, [8]=Exp, [9]=Prio, [10]=Type, [11]=Tool, [12]=Stat
        all_rows.append(("Mobile (Appium)", row))

    # Read Selenium (Rows 3 to 302)
    for row in ws_sel.iter_rows(min_row=3, max_row=302, values_only=True):
        all_rows.append(("Web (Selenium)", row))

    # Populate Rows (Row 6 to 605)
    r = 6
    for idx, (platform, tc_vals) in enumerate(all_rows, 1):
        tc_id, m_name, submod, title, desc, pre, steps, data, exp, prio, ttype, tool, stat = tc_vals
        ws.row_dimensions[r].height = 24
        bg_row = c_bg_subtle if r % 2 == 0 else "FFFFFF"

        row_cells = [
            ("A", idx, "center", True),
            ("B", tc_id, "center", True),
            ("C", platform, "center", True),
            ("D", m_name, "left", False),
            ("E", submod, "left", False),
            ("F", title, "left", True),
            ("G", desc, "left", False),
            ("H", pre, "left", False),
            ("I", steps, "left", False),
            ("J", data, "left", False),
            ("K", exp, "left", False),
            ("L", prio, "center", True),
            ("M", ttype, "center", False),
            ("N", tool, "center", False),
            ("O", stat, "center", True),
        ]

        for col_char, val, align, is_bold in row_cells:
            cell = ws[f"{col_char}{r}"]
            cell.value = val
            cell.font = Font(name="Segoe UI", size=8.5, bold=is_bold, color=c_text_body)
            cell.fill = PatternFill(start_color=bg_row, end_color=bg_row, fill_type="solid")
            cell.alignment = Alignment(horizontal=align, vertical="center", wrap_text=(col_char in ["F", "G", "I", "K"]))
            cell.border = thin_border_light

            # Platform styling
            if col_char == "C":
                if platform == "Mobile (Appium)":
                    cell.fill = PatternFill(start_color=c_primary_light, end_color=c_primary_light, fill_type="solid")
                    cell.font = Font(name="Segoe UI", size=8.5, bold=True, color=c_primary_deep)
                else:
                    cell.fill = PatternFill(start_color=c_teal_light, end_color=c_teal_light, fill_type="solid")
                    cell.font = Font(name="Segoe UI", size=8.5, bold=True, color=c_teal_dark)

            # Priority badges
            if col_char == "L":
                if val == "Critical":
                    cell.fill = PatternFill(start_color=c_crit_bg, end_color=c_crit_bg, fill_type="solid")
                    cell.font = Font(name="Segoe UI", size=8.5, bold=True, color=c_crit_fg)
                elif val == "High":
                    cell.fill = PatternFill(start_color=c_high_bg, end_color=c_high_bg, fill_type="solid")
                    cell.font = Font(name="Segoe UI", size=8.5, bold=True, color=c_high_fg)
                elif val == "Medium":
                    cell.fill = PatternFill(start_color=c_med_bg, end_color=c_med_bg, fill_type="solid")
                    cell.font = Font(name="Segoe UI", size=8.5, bold=False, color=c_med_fg)
                elif val == "Low":
                    cell.fill = PatternFill(start_color=c_low_bg, end_color=c_low_bg, fill_type="solid")
                    cell.font = Font(name="Segoe UI", size=8.5, bold=False, color=c_low_fg)

            # Status badges
            if col_char == "O":
                if val == "Passed":
                    cell.fill = PatternFill(start_color=c_pass_bg, end_color=c_pass_bg, fill_type="solid")
                    cell.font = Font(name="Segoe UI", size=8.5, bold=True, color=c_pass_fg)
                elif val == "Automated":
                    cell.fill = PatternFill(start_color=c_auto_bg, end_color=c_auto_bg, fill_type="solid")
                    cell.font = Font(name="Segoe UI", size=8.5, bold=True, color=c_auto_fg)
                elif val == "In Progress":
                    cell.fill = PatternFill(start_color=c_prog_bg, end_color=c_prog_bg, fill_type="solid")
                    cell.font = Font(name="Segoe UI", size=8.5, bold=True, color=c_prog_fg)

        r += 1

    # Freeze panes below Row 5 so the header and KPI banner stay visible
    ws.freeze_panes = "A6"

    # AutoFilter on Row 5 covering all 600 data rows
    ws.auto_filter.ref = f"A5:O{r-1}"

    # Save to file
    out_file = r"c:\Users\srinu\healthcare_management\healthcare_ehr_all_test_cases_600.xlsx"
    wb.save(out_file)
    print(f"Master single-sheet workbook saved to: {out_file}")

    docs_file = r"c:\Users\srinu\healthcare_management\docs\testing\healthcare_ehr_all_test_cases_600.xlsx"
    wb.save(docs_file)
    print(f"Copy also saved to: {docs_file}")

if __name__ == "__main__":
    build_single_sheet_excel()
